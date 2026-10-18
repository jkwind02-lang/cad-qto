"""Portable two-sheet delivery: facts, empty/zero, navigation and same-source HTML."""
import copy
import hashlib
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / 'skills/_qto-shared-evidence/scripts'
sys.path.insert(0, str(SCRIPTS))
from build_qto_delivery_excel import build, create_workbook, verify_workbook, comparison_for, SUMMARY, DETAIL, MAX_TEXT
from build_qto_review_html import build as build_html
from qto_review_contract import normalize_view
from review_delivery_fixture import view, scope
from review_comparison_fixture import view as comparison_view


def write_input(tmp_path, raw, name='input.json'):
    target = tmp_path / name
    target.write_text(json.dumps(raw, ensure_ascii=False), encoding='utf-8')
    return target


def test_opt_in_delivery_without_items_and_actual_html(tmp_path):
    raw = view(with_scene=False)
    before = copy.deepcopy(raw)
    data = write_input(tmp_path, raw)
    html = tmp_path / 'review.html'
    build_html(data, html, feedback_url='')
    report = build(data, tmp_path / '两表.xlsx', html_path=html)
    assert report['status'] == 'PASS'
    assert report['same_source_html']['pass'] is True
    assert report['states'] == {'current': 2, 'reference': 2, 'historical': 0, 'retired': 1}
    assert report['current_null_count'] == 3
    assert report['hyperlinks_checked'] >= 10
    assert raw == before
    wb = load_workbook(tmp_path / '两表.xlsx')
    assert wb.sheetnames == [SUMMARY, DETAIL]
    assert all(w.sheet_state == 'visible' for w in wb)
    assert all(not w.row_dimensions[n].hidden for w in wb for n in w.row_dimensions)
    wb.close()


def test_long_evidence_not_truncated_and_text_never_formula(tmp_path):
    raw = view()
    long_basis = '=1+1\n' + '根据合成记录核对长度😀。' * 6000
    raw['delivery']['quantity_rows'][0]['basis'] = long_basis
    raw['delivery']['quantity_rows'][0]['title'] = '=未执行的来源文字'
    data = write_input(tmp_path, raw)
    report = build(data, tmp_path / '长依据.xlsx')
    assert report['status'] == 'PASS'
    assert report['formula_count'] == 0
    wb = load_workbook(tmp_path / '长依据.xlsx')
    ws = wb[DETAIL]
    assert ws['J6'].data_type == 's'
    pieces = [ws.cell(6 + i, 10).value or '' for i in range((len(long_basis) + MAX_TEXT - 1) // MAX_TEXT)]
    assert ''.join(pieces) == long_basis
    assert all(len(x.encode('utf-16-le')) // 2 < 32767 for x in pieces)
    wb.close()


def test_existing_outputs_and_source_path_protected(tmp_path):
    data = write_input(tmp_path, view())
    with pytest.raises(ValueError, match='不能覆盖'):
        build(data, data)
    output = tmp_path / 'protected.xlsx'
    output.write_bytes(b'user edits')
    with pytest.raises(ValueError, match='已存在'):
        build(data, output)
    assert output.read_bytes() == b'user edits'


def test_no_delivery_uses_explicit_conversion_error(tmp_path):
    raw = view()
    raw.pop('delivery')
    data = write_input(tmp_path, raw)
    with pytest.raises(ValueError, match='convert_qto_review_delivery'):
        build(data, tmp_path / '不生成.xlsx')
    assert not (tmp_path / '不生成.xlsx').exists()


def test_foreign_html_rejected_before_output(tmp_path):
    raw = view()
    data = write_input(tmp_path, raw)
    foreign = copy.deepcopy(raw)
    foreign['delivery']['quantity_rows'][0]['basis'] = '另一版本的依据'
    foreign_input = write_input(tmp_path, foreign, 'foreign.json')
    html = tmp_path / 'foreign.html'
    build_html(foreign_input, html, feedback_url='')
    with pytest.raises(ValueError, match='不是同一'):
        build(data, tmp_path / '不生成.xlsx', html_path=html)
    assert not (tmp_path / '不生成.xlsx').exists()


def test_tampering_is_detected_in_displayed_quantities(tmp_path):
    normalized = normalize_view(view())
    wb, model, positions = create_workbook(normalized)
    wb[DETAIL]['G6'] = 999
    output = tmp_path / 'tampered.xlsx'
    wb.save(output)
    with pytest.raises(ValueError, match='明细不一致'):
        verify_workbook(output, normalized, model, positions)


def test_proven_current_zero_stays_numeric_and_real_decisions_remain_data(tmp_path):
    raw = view(decisions=True)
    raw['delivery']['quantity_rows'][1]['quantity'] = 0
    data = write_input(tmp_path, raw)
    result = build(data, tmp_path / '有据零.xlsx')
    assert result['status'] == 'PASS'
    wb = load_workbook(tmp_path / '有据零.xlsx')
    assert wb[DETAIL]['G7'].value == 0 and wb[DETAIL]['G7'].data_type == 'n'
    assert wb[SUMMARY]['F7'].value == 0
    wb.close()


def test_scope_only_reference_and_migration_notice_never_claim_verified_inclusion(tmp_path):
    raw = view()
    reference = raw['delivery']['quantity_rows'][2]['reference']
    reference.update(association_state='scope_only', relationship='unknown', current_row_ids=[],
                     source_claimed_relationship='includes_current', scope_statement='原记录称全支路',
                     association_reason='未逐明细确认包含范围')
    notice = '来源已采用记录迁移，本次未重算；未重新证实几何或规则输入，不代表正式核准。'
    raw['delivery']['projection_provenance'] = {'notice': notice,
        'input_sha256': hashlib.sha256(json.dumps(raw, ensure_ascii=False).encode('utf-8')).hexdigest(),
        'verification_scope': 'frozen_record_identity_only'}
    data = write_input(tmp_path, raw)
    build(data, tmp_path / '范围关联.xlsx')
    wb = load_workbook(tmp_path / '范围关联.xlsx')
    # The shared normalizer supplies the fixed boundary, overriding free notes.
    assert normalize_view(raw)['delivery']['projection_provenance']['notice'] in wb[SUMMARY]['A4'].value
    assert '逐明细包含范围未证' in wb[DETAIL]['L8'].value
    assert wb[DETAIL]['G8'].value is None
    wb.close()


def test_only_formally_comparable_matching_source_delta_is_consumed():
    normalized = normalize_view(comparison_view(with_scene=False))
    item = {'id': 'I2', 'unit': '个', 'quantity': 10, 'original_quantity': 11}
    delta, rate, reason = comparison_for(normalized, item)
    assert delta == -1 and rate == -0.0909 and '可比' in reason
    item['quantity'] = None
    assert comparison_for(normalized, item)[:2] == (None, None)
    multiple = {'id': 'I1', 'unit': 'm', 'quantity': 35, 'original_quantity': 34}
    assert comparison_for(normalized, multiple)[:2] == (None, None)


def test_valid_decimal_serialization_uses_same_tolerance_in_detail_and_summary(tmp_path):
    raw = view()
    raw['delivery']['quantity_rows'][0]['quantity'] = 12.188465199696114
    data = write_input(tmp_path, raw)
    report = build(data, tmp_path / '合法小数.xlsx')
    assert report['status'] == 'PASS'
    assert report['row_checks']['quantity_rows'][0]['quantity'] == 12.188465199696114


def test_adopted_record_has_fixed_visible_identity_only_boundary_without_notice(tmp_path):
    raw = view()
    source_sha = hashlib.sha256(json.dumps(raw, ensure_ascii=False).encode('utf-8')).hexdigest()
    raw['delivery']['projection_provenance'] = {'input_sha256': source_sha,
        'verification_scope': 'frozen_record_identity_only'}
    for row in raw['delivery']['quantity_rows'][:2]:
        row['evidence'] = {'state': 'verified', 'method': 'adopted_record', 'new_measurement_performed': False,
            'source_refs': [], 'input_facts': [], 'provenance': {'input_sha256': source_sha,
            'record_id': row['id'], 'record_locator': '/frozen/' + row['id'], 'role_path': 'role',
            'role_value': 'current', 'verification_scope': 'frozen_record_identity_only'}}
    data = write_input(tmp_path, raw)
    build(data, tmp_path / '记录迁移.xlsx')
    wb = load_workbook(tmp_path / '记录迁移.xlsx')
    for sheet in wb:
        assert '本次未重算' in sheet['A4'].value and '仅核对原记录身份' in sheet['A4'].value
    assert '本次未重算' in wb[DETAIL]['K6'].value
    assert '证据已核实' not in wb[DETAIL]['K6'].value
    assert 'adopted_record' not in wb[DETAIL]['K6'].value
    wb.close()


def test_declared_items_preserve_unknown_explicit_zero_and_mixed_units(tmp_path):
    raw = view(with_scene=False)
    rows = raw['delivery']['quantity_rows']
    items = []
    for rid, title, qty, unit, ids, original in [
            ('ITEM-A', '管道', 12, 'm', ['M1'], 10),
            ('ITEM-B', '泵组', 2, '台', ['M2'], 0),
            ('ITEM-C', '待补竖向长度', None, 'm', [], None)]:
        source = next((r for r in rows if r['id'] in ids), rows[3])
        items.append({'id': rid, 'title': title, 'quantity': qty, 'original_quantity': original,
                      'unit': unit, 'scope': copy.deepcopy(source['scope']), 'coverage': copy.deepcopy(source['coverage']),
                      'current_row_ids': ids, 'reference_row_ids': ['R1', 'R2'] if rid == 'ITEM-A' else [],
                      'historical_row_ids': ['H1'] if rid == 'ITEM-C' else [], 'dimensions': {}})
    raw['delivery']['items'] = items
    items[2]['coverage']['state'] = 'unknown'
    items[0]['original_unit'] = '组'
    data = write_input(tmp_path, raw)
    report = build(data, tmp_path / '有条目.xlsx')
    assert report['item_count'] == 3
    wb = load_workbook(tmp_path / '有条目.xlsx')
    ws = wb[SUMMARY]
    assert ws['G7'].value == 0
    assert ws['F8'].value is None and ws['G8'].value is None
    assert ws['I6'].value is None and ws['J6'].value is None
    assert ws['H6'].value == '组'
    assert [ws.cell(i, 5).value for i in (6, 7, 8)] == ['m', '台', 'm']
    assert not any(c.data_type == 'f' for sheet in wb for row in sheet for c in row)
    wb.close()
