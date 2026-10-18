"""Optional two-sheet Excel delivery from the same normalized view as review HTML.

No recognition, estimation, approval, mutation of input, or automatic HTML output.
Requires openpyxl. All records remain visible; unknown quantities remain blank.
"""
from __future__ import annotations

import argparse
import base64
import copy
import gzip
import hashlib
import json
import math
import re
import tempfile
from pathlib import Path

from qto_review_contract import normalize_view, unpack_payload


SUMMARY = '工程量汇总'
DETAIL = '工程量明细'
HEADER_ROW = 5
START_ROW = 6
MAX_TEXT = 15000  # Also safe for UTF-16 surrogate pairs under Excel's 32767 limit.
SCHEMA = 'qto-delivery-excel-check/1.0'
RECORDED_BOUNDARY = '来源已采用记录迁移，本次未重算；仅核对原记录身份，未重新证实几何或规则输入，不代表正式核准。'
STATE_LABELS = {'current': '本次主量', 'reference': '参考待核（不相加）',
                'historical': '历史留档（不采用）', 'retired': '退役留档（不采用）'}
SUMMARY_HEADERS = ['条目编号', '名称', '规格及维度', '适用及排除范围', '单位', '本次采用量',
                   '原量', '原量单位', '可比差额', '偏差比例', '比较资格及具体原因', '覆盖情况',
                   '主量明细', '参考明细', '历史明细', '完整业务记录（JSON）']
DETAIL_HEADERS = ['汇总编号', '明细编号', '数量状态', '适用及排除范围', '对象名称', '单位',
                  '本次主量', '参考量', '历史及退役量', '计算式及依据', '证据及定位',
                  '参考关系及适用前提', '覆盖情况', '成员编号及证据成员', '返回汇总', '完整业务记录（JSON）']


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)


def digest(value):
    return hashlib.sha256(canonical(value).encode('utf-8')).hexdigest()


def equal_numeric(a, b):
    return a == b or (isinstance(a, (int, float)) and isinstance(b, (int, float))
                      and math.isclose(a, b, rel_tol=1e-14, abs_tol=1e-14))


def html_delivery_view(path):
    text = Path(path).read_text(encoding='utf-8-sig')
    match = re.search(r'<script\s+id="payload"\s+type="application/octet-stream">([\s\S]*?)</script>', text)
    if not match:
        raise ValueError('HTML未包含共享生成器payload数据块；不执行其中脚本')
    try:
        return unpack_payload(gzip.decompress(base64.b64decode(match.group(1))))
    except (ValueError, OSError) as error:
        raise ValueError('HTML共享数据块不能解包') from error


def delivery_data(view):
    # Same public helper as the HTML delivery contract; compatible with the first
    # opt-in delivery/1.0 release which did not yet export the helper.
    try:
        from qto_review_delivery import canonical_delivery_data
    except ImportError:
        return {k: copy.deepcopy(v) for k, v in view['delivery'].items()
                if k not in ('generation_check', 'comparison_statistics', 'coverage_check')}
    return canonical_delivery_data(view)


def shown(value):
    if value is None:
        return ''
    if isinstance(value, str):
        return value
    return canonical(value)


def scope_text(value):
    value = value or {}
    return '适用：' + '、'.join(value.get('included', [])) + '\n排除：' + ('、'.join(value.get('excluded', [])) or '未列排除范围')


def coverage_text(value):
    value = value or {}
    state = {'complete': '完整覆盖', 'partial': '局部覆盖', 'unknown': '未形成采用量'}.get(value.get('state'), '未声明')
    return state + '\n已算：' + '、'.join(value.get('computed_scope', [])) + '\n未算：' + '、'.join(value.get('uncomputed_scope', []))


def evidence_text(value):
    value = value or {}
    if value.get('method') == 'adopted_record':
        lines = ['来源采用记录迁移；本次未重算；仅核对原记录身份',
                 '未重新证实几何或规则输入，不代表正式核准。']
        provenance = value.get('provenance', {})
        if provenance.get('record_locator'):
            lines.append('原记录定位：' + provenance['record_locator'])
        lines.extend(ref['source_id'] + '：' + ref['locator'] for ref in value.get('source_refs', []))
        for fact in value.get('input_facts', []):
            lines.append('原记录事实：' + fact.get('statement', ''))
            lines.extend(ref['source_id'] + '：' + ref['locator'] for ref in fact.get('source_refs', []))
        return '\n'.join(lines)
    state = {'verified': '证据已核实', 'pending': '证据待核'}.get(value.get('state'), '原记录证据')
    method = {'measured': '实测', 'rule': '规则推算', 'recorded': '来源采用记录'}.get(value.get('method'), shown(value.get('method')))
    lines = [state + ('；' + method if method else '')]
    lines.extend(ref['source_id'] + '：' + ref['locator'] for ref in value.get('source_refs', []))
    for fact in value.get('input_facts', []):
        lines.append('输入事实：' + fact.get('statement', '') + '（' + {'verified': '已核实', 'pending': '待核'}.get(fact.get('state'), '原记录') + '）')
        lines.extend(ref['source_id'] + '：' + ref['locator'] for ref in fact.get('source_refs', []))
    return '\n'.join(lines)


def items_for(view):
    d = view['delivery']
    if 'items' in d:
        return copy.deepcopy(d['items'])
    # Opt-in delivery/1.0 compatibility: one row per item. A reference or history
    # row never becomes current merely because an item list was not supplied.
    result = []
    for row in d['quantity_rows']:
        kind = 'current' if row['state'] == 'current' else 'reference' if row['state'] == 'reference' else 'historical'
        item = {'id': row['id'], 'title': row['title'], 'unit': row['unit'],
                'quantity': row['quantity'] if kind == 'current' else None,
                'original_quantity': None, 'scope': copy.deepcopy(row['scope']),
                'coverage': copy.deepcopy(row['coverage']), 'dimensions': {},
                'current_row_ids': [], 'reference_row_ids': [], 'historical_row_ids': []}
        item[kind + '_row_ids'] = [row['id']]
        result.append(item)
    return result


def comparison_for(view, item):
    """Only consume professional comparison results; never invent eligibility."""
    source = next((x for x in (view.get('comparison') or {}).get('items', []) if x['id'] == item.get('comparison_item_id', item['id'])), None)
    segments = source.get('third_parties', []) if source else []
    if not segments:
        return None, None, '未提供正式可比条件及差额；不以同单位自动判为可比'
    reasons = []
    state_labels = {'comparable': '可比', 'not_comparable': '不可比', 'unknown': '待核可比性',
                    'original_unknown': '原量未知', 'current_unknown': '本次量未形成',
                    'unit_mismatch': '单位不同', 'not_in_scope': '本期不计'}
    # One visible numeric delta only when one normalized source segment exactly
    # matches this item's current/original/unit. Multiple sources are not summed.
    for seg in segments:
        reasons.append(seg['source_id'] + '：' + state_labels.get(seg.get('comparison_state'), '待核可比性') + '；' +
                       (seg.get('comparability') or {}).get('reason', '') + '；' + shown(seg.get('conclusion')))
    if len(segments) == 1:
        seg = segments[0]
        if (seg.get('comparison_state') == 'comparable' and seg.get('unit') == item['unit']
                and seg.get('quantity') == item.get('original_quantity')
                and source['ours'].get('total') == item.get('quantity')):
            diff = seg.get('difference') or {}
            return diff.get('value'), diff.get('rate'), '\n'.join(reasons)
    if len(segments) > 1:
        reasons.append('多个独立来源段不跨来源相加；逐段结果完整保留在业务记录')
    else:
        reasons.append('当前采用、原量或单位与比较段不一致时，不沿用旧差额')
    return None, None, '\n'.join(reasons)


def anchor(prefix, identifier):
    return prefix + hashlib.sha256(identifier.encode('utf-8')).hexdigest()[:24]


def model_for(view):
    items = items_for(view)
    rows = copy.deepcopy(view['delivery']['quantity_rows'])
    owner = {}
    byrow = {r['id']: r for r in rows}
    for item in items:
        for field in ('current_row_ids', 'reference_row_ids', 'historical_row_ids'):
            for rid in item.get(field, []):
                if rid not in byrow:
                    raise ValueError('汇总引用不存在的明细：' + rid)
                owner.setdefault(rid, []).append(item['id'])
    # No current/reference/history row may vanish from the two-sheet output.
    return {'items': items, 'rows': rows, 'owner': owner, 'delivery': delivery_data(view),
            'comparison': copy.deepcopy(view.get('comparison', {})),
            'sources': copy.deepcopy(view['sources']),
            'project_id': view['project_id'], 'review_id': view['review_id'], 'revision': view['revision'],
            'source_sha256': view['source_sha256'], 'result_fingerprint': view['result_fingerprint']}


def _put(cell, value):
    if isinstance(value, str):
        # Source-controlled text such as '=CMD(...)' must never become a formula.
        cell.value = value
        cell.data_type = 's'
    else:
        cell.value = value


def _append(ws, values):
    """Split every oversized text without dropping a character, even raw JSON."""
    start = max(ws.max_row + 1, START_ROW)
    parts = [([v[i:i + MAX_TEXT] for i in range(0, len(v), MAX_TEXT)] or [''])
             if isinstance(v, str) else [v] for v in values]
    count = max(len(p) for p in parts)
    if start + count - 1 > 1048576:
        raise ValueError('完整数据超出单张Excel工作表行数上限；请按业务范围拆分输入')
    for i in range(count):
        for col, part in enumerate(parts, 1):
            _put(ws.cell(start + i, col), part[i] if i < len(part) else None)
        if i:
            if len(parts[0]) == 1:
                _put(ws.cell(start + i, 1), values[0])
            if len(parts[1]) == 1:
                _put(ws.cell(start + i, 2), '续文' + str(i))
    return {'start': start, 'count': count}


def _join(ws, position, column):
    return ''.join(str(ws.cell(position['start'] + n, column).value or '') for n in range(position['count']))


def _setup_sheet(ws, title, headers, view, content_hash):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    ws.sheet_view.showGridLines = False
    _put(ws['A2'], title)
    ws['A2'].font = Font(name='Microsoft YaHei', size=15, bold=True, color='17365D')
    _put(ws['A3'], view['title'] + '；版本：' + view['revision'])
    provenance = view['delivery'].get('projection_provenance', {}).get('notice', '')
    if any(r.get('evidence', {}).get('method') == 'adopted_record' for r in view['delivery']['quantity_rows']):
        provenance = RECORDED_BOUNDARY
    _put(ws['A4'], (provenance + '\n' if provenance else '') + '未知数量留空，有据零保留0。参考和历史不计入主量；不同单位不求总量。')
    if provenance:
        ws.row_dimensions[4].height = 32
        ws['A4'].alignment = Alignment(wrap_text=False, vertical='top')
    _put(ws['J2'], '同源数据指纹')
    _put(ws['K2'], content_hash)
    _put(ws['J3'], '结果版本指纹')
    _put(ws['K3'], view['result_fingerprint'])
    _put(ws['J4'], view['delivery']['adoption_note'])
    for address in ('A3', 'A4', 'J2', 'K2', 'J3', 'K3', 'J4'):
        ws[address].font = Font(name='Microsoft YaHei', size=10, color='444444')
    for col, label in enumerate(headers, 1):
        cell = ws.cell(HEADER_ROW, col, label)
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(name='Microsoft YaHei', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(bottom=Side(style='thin', color='B7C9E2'))
    ws.row_dimensions[HEADER_ROW].height = 32
    ws.freeze_panes = 'F6'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:5'


def _style_data(ws, widths, number_columns):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    for n, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(n)].width = width
    for row in ws.iter_rows(min_row=START_ROW, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name='Microsoft YaHei', size=10, color='222222')
            cell.alignment = Alignment(vertical='top', wrap_text=True,
                                       horizontal='right' if cell.column in number_columns else 'left')
            if cell.column in number_columns:
                cell.number_format = '#,##0' if isinstance(cell.value, (int, float)) and cell.value == int(cell.value) else '#,##0.###############'
        # The complete value stays inspectable in each cell and its raw record.
        # A restrained row height preserves usable navigation for large inputs.
        visible = [c.value for c in row[:-1] if c.value is not None]
        height = min(90, max(30, 15 * max((str(x).count('\n') + 1 for x in visible), default=1)))
        ws.row_dimensions[row[0].row].height = height
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill('solid', fgColor='F3F6FA')
    ws.auto_filter.ref = f'A{HEADER_ROW}:{get_column_letter(len(widths))}{max(HEADER_ROW, ws.max_row)}'


def _name(wb, name, sheet, row):
    from openpyxl.workbook.defined_name import DefinedName
    dn = DefinedName(name, attr_text=f"'{sheet}'!$A${row}")
    if hasattr(wb.defined_names, 'add'):
        wb.defined_names.add(dn)
    else:
        wb.defined_names.append(dn)


def _link(cell, name):
    from openpyxl.styles import Font
    from openpyxl.worksheet.hyperlink import Hyperlink
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=name, display=str(cell.value or '查看'))
    cell.font = Font(name='Microsoft YaHei', size=10, color='0563C1', underline='single')


def create_workbook(view):
    try:
        from openpyxl import Workbook
    except ImportError as error:
        raise ValueError('Excel入口需要openpyxl，请在本Python环境安装：python -m pip install openpyxl') from error
    model = model_for(view)
    content_hash = digest(model['delivery'])
    wb = Workbook()
    summary = wb.active
    summary.title = SUMMARY
    detail = wb.create_sheet(DETAIL)
    wb.properties.title = view['title']
    wb.properties.subject = '工程量汇总与明细，同一规范化交付视图'
    wb.properties.identifier = content_hash
    wb.properties.keywords = 'result_fingerprint=' + view['result_fingerprint']
    _setup_sheet(summary, SUMMARY, SUMMARY_HEADERS, view, content_hash)
    _setup_sheet(detail, DETAIL, DETAIL_HEADERS, view, content_hash)
    positions = {'items': {}, 'rows': {}, 'metadata': []}
    for item in model['items']:
        delta, rate, reason = comparison_for(view, item)
        comp = next((x for x in model['comparison'].get('items', []) if x['id'] == item.get('comparison_item_id', item['id'])), None)
        # Includes the full normalized source comparison, not a recomputed delta.
        raw = {'item': item, 'comparison': comp}
        dimensions = '\n'.join(k + '：' + shown(v) for k, v in item.get('dimensions', {}).items())
        values = [item['id'], item['title'], dimensions, scope_text(item.get('scope')),
                  item['unit'], item.get('quantity'), item.get('original_quantity'), item.get('original_unit', item['unit']), delta, rate, reason,
                  coverage_text(item.get('coverage')), '\n'.join(item.get('current_row_ids', [])),
                  '\n'.join(item.get('reference_row_ids', [])), '\n'.join(item.get('historical_row_ids', [])), canonical(raw)]
        positions['items'][item['id']] = _append(summary, values)
    for row in model['rows']:
        ref = row.get('reference', {})
        relationship = {'includes_current': '包含所列主量', 'excludes_current': '不含所列主量', 'unknown': '是否含主量待核'}.get(ref.get('relationship'), '')
        ref_text = (relationship + '\n关联主量：' + '、'.join(ref.get('current_row_ids', [])) +
                    '\n前提：' + '；'.join(ref.get('prerequisites', []))) if ref else row.get('history_reason', '')
        if ref.get('association_state') == 'scope_only':
            ref_text = ('原记录关系声明；逐明细包含范围未证，不得相加。\n原关系记录：' + ref.get('source_claimed_relationship', ref.get('relationship', '')) +
                        '\n范围说明：' + shown(ref.get('scope_statement')) + '\n关联边界：' + shown(ref.get('association_reason')) +
                        '\n前提：' + '；'.join(ref.get('prerequisites', [])))
        members = '数量成员：' + '、'.join(row.get('member_keys', [])) + '\n仅证据成员：' + '、'.join(row.get('evidence_member_keys', []))
        owners = model['owner'].get(row['id'], [])
        state = row['state']
        values = ['\n'.join(owners), row['id'], STATE_LABELS[state], scope_text(row.get('scope')), row['title'], row['unit'],
                  row['quantity'] if state == 'current' else None,
                  row['quantity'] if state == 'reference' else None,
                  row['quantity'] if state in ('historical', 'retired') else None,
                  row['basis'], evidence_text(row.get('evidence', {})), ref_text, coverage_text(row.get('coverage')),
                  members, '\n'.join(owners) if owners else '无汇总归属（完整保留）', canonical(row)]
        positions['rows'][row['id']] = _append(detail, values)
    def append_missing(group):
        needs = group.get('minimum_required', [])
        explanation = '\n'.join('缺件：' + need['description'] + '\n提供方：' + need['provider'] +
                                '\n影响：' + need['impact'] + '\n回算动作：' + need['recompute_action'] for need in needs)
        p = _append(detail, ['资料缺口', group['id'], '缺资料（不填0）', scope_text(group['scope']), group['title'], '',
                         None, None, None, explanation or '缺件详见以下独立子项', shown(group.get('verified_facts', [])),
                         '', '', '', '汇总首页', canonical(group)])
        positions.setdefault('missing', []).append(p)
        for child in group.get('children', []):
            append_missing(child)
    for group in model['delivery'].get('missing_groups', []):
        append_missing(group)
    # Keep uncensored missing groups, decision targets, extra fields, source
    # membership and coverage available in the same second sheet. No hidden tab.
    # Chunk the entire canonical delivery record for exact fingerprint readback.
    metadata = {'delivery': model['delivery'], 'comparison': model['comparison'], 'sources': model['sources'],
                'identity': {k: model[k] for k in ('project_id', 'review_id', 'revision', 'source_sha256', 'result_fingerprint')}}
    metadata_text = canonical(metadata)
    pos = _append(detail, ['完整交付依据', 'DELIVERY-SOURCE', '原始业务留档', '', '缺件、取舍、成员与比较条件完整记录', '',
                           None, None, None, '完整业务记录列连续保存全部规范化交付数据。', '', '', '', '', '返回汇总首页', metadata_text])
    positions['metadata'].append(pos)
    _style_data(summary, [20, 34, 30, 38, 10, 15, 15, 12, 15, 15, 48, 34, 24, 24, 24, 70], {6, 7, 9, 10})
    _style_data(detail, [20, 24, 25, 38, 34, 10, 15, 15, 15, 70, 65, 48, 34, 45, 20, 70], {7, 8, 9})
    from openpyxl.styles import Alignment
    for cells in summary.iter_rows(min_row=START_ROW, min_col=8, max_col=8):
        cells[0].alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
    for item in model['items']:
        position = positions['items'][item['id']]['start']
        _name(wb, anchor('item_', item['id']), SUMMARY, position)
        for col, key in ((13, 'current_row_ids'), (14, 'reference_row_ids'), (15, 'historical_row_ids')):
            ids = item.get(key, [])
            if ids:
                _link(summary.cell(position, col), anchor('row_', ids[0]))
        summary.cell(position, 10).number_format = '0.00%'
    for row in model['rows']:
        position = positions['rows'][row['id']]['start']
        _name(wb, anchor('row_', row['id']), DETAIL, position)
        owners = model['owner'].get(row['id'], [])
        if owners:
            _link(detail.cell(position, 15), anchor('item_', owners[0]))
    _name(wb, 'summary_home', SUMMARY, 2)
    _link(detail.cell(pos['start'], 15), 'summary_home')
    for missing_position in positions.get('missing', []):
        _link(detail.cell(missing_position['start'], 15), 'summary_home')
    _link(summary['N2'], 'summary_home')
    _put(summary['N2'], '汇总首页')
    return wb, model, positions


def verify_workbook(path, view, model, positions, html_path=None):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False)
    errors = []
    checks = {'items': [], 'quantity_rows': []}
    if wb.sheetnames != [SUMMARY, DETAIL]:
        errors.append('主表名称或顺序不一致')
    summary, detail = wb[SUMMARY], wb[DETAIL]
    def check_name(name, sheet, row):
        definition = wb.defined_names.get(name)
        return definition is not None and definition.attr_text == f"'{sheet}'!$A${row}"
    for item in model['items']:
        p = positions['items'][item['id']]
        n = p['start']
        record = json.loads(_join(summary, p, 16))
        expected = [item['id'], item['title'], item['unit'], item.get('quantity'), item.get('original_quantity'), item.get('original_unit', item['unit'])]
        actual = [summary.cell(n, c).value for c in (1, 2, 5, 6, 7, 8)]
        delta, rate, _ = comparison_for(view, item)
        ok = (all(equal_numeric(a, b) for a, b in zip(expected, actual)) and record['item'] == item
              and equal_numeric(summary.cell(n, 9).value, delta) and equal_numeric(summary.cell(n, 10).value, rate))
        checks['items'].append({'id': item['id'], 'unit': item['unit'], 'quantity': item.get('quantity'),
                                'original_quantity': item.get('original_quantity'), 'difference': delta, 'rate': rate, 'pass': ok})
        if not ok:
            errors.append('汇总不一致：' + item['id'])
        if not check_name(anchor('item_', item['id']), SUMMARY, n):
            errors.append('汇总稳定定位不一致：' + item['id'])
        for col, field in ((13, 'current_row_ids'), (14, 'reference_row_ids'), (15, 'historical_row_ids')):
            ids = item.get(field, [])
            link = summary.cell(n, col).hyperlink
            if bool(ids) != bool(link) or ids and link.location != anchor('row_', ids[0]):
                errors.append('汇总明细跳转不一致：' + item['id'])
    for row in model['rows']:
        p = positions['rows'][row['id']]
        n = p['start']
        record = json.loads(_join(detail, p, 16))
        qty_col = 7 if row['state'] == 'current' else 8 if row['state'] == 'reference' else 9
        actual = detail.cell(n, qty_col).value
        other_blank = all(detail.cell(n, col).value is None for col in (7, 8, 9) if col != qty_col)
        member_text = '数量成员：' + '、'.join(row.get('member_keys', [])) + '\n仅证据成员：' + '、'.join(row.get('evidence_member_keys', []))
        ok = (record == row and detail.cell(n, 6).value == row['unit'] and equal_numeric(actual, row['quantity']) and other_blank
              and detail.cell(n, 3).value == STATE_LABELS[row['state']]
              and _join(detail, p, 10) == row['basis'] and _join(detail, p, 14) == member_text)
        checks['quantity_rows'].append({'id': row['id'], 'state': row['state'], 'unit': row['unit'],
                                       'quantity': row['quantity'], 'member_keys': row['member_keys'], 'pass': ok})
        if not ok:
            errors.append('明细不一致：' + row['id'])
        if not check_name(anchor('row_', row['id']), DETAIL, n):
            errors.append('明细稳定定位不一致：' + row['id'])
        owners = model['owner'].get(row['id'], [])
        link = detail.cell(n, 15).hyperlink
        if bool(owners) != bool(link) or owners and link.location != anchor('item_', owners[0]):
            errors.append('明细返回跳转不一致：' + row['id'])
    metadata = json.loads(_join(detail, positions['metadata'][0], 16))
    content_hash = digest(delivery_data(view))
    if metadata['delivery'] != model['delivery'] or digest(metadata['delivery']) != content_hash or wb.properties.identifier != content_hash:
        errors.append('完整交付记录或同源指纹不一致')
    if metadata['comparison'] != model['comparison']:
        errors.append('完整比较条件不一致')
    if metadata['sources'] != model['sources']:
        errors.append('来源编号、名称或身份不一致')
    hyperlink_count = 0
    for ws in wb:
        if ws.sheet_state != 'visible' or ws.sheet_view.showGridLines is not False or ws.freeze_panes != 'F6':
            errors.append(ws.title + '显示或冻结设置异常')
        for cells in ws.iter_rows():
            for cell in cells:
                if cell.data_type == 'f':
                    errors.append('源文本被写成公式：' + ws.title + '!' + cell.coordinate)
                if cell.hyperlink:
                    hyperlink_count += 1
                    if cell.hyperlink.location not in wb.defined_names:
                        errors.append('跳转目标不存在：' + str(cell.hyperlink.location))
    html_check = None
    if html_path:
        # Read the actual embedded normalized view using the HTML extractor used
        # by feedback import. Byte equality of the business record is required.
        html_view = html_delivery_view(html_path)
        if not isinstance(html_view, dict) or 'delivery' not in html_view:
            raise ValueError('指定HTML未包含delivery视图')
        html_check = {'path': str(Path(html_path).resolve()),
                      'delivery_sha256': digest(delivery_data(html_view)),
                      'source_sha256': html_view.get('source_sha256'),
                      'result_fingerprint': html_view.get('result_fingerprint')}
        html_check['pass'] = (html_check['delivery_sha256'] == content_hash
                              and html_check['source_sha256'] == view['source_sha256']
                              and html_check['result_fingerprint'] == view['result_fingerprint'])
        if not html_check['pass']:
            errors.append('HTML与Excel不是同一规范化交付数据或结果版本')
    report = {'schema': SCHEMA, 'status': 'PASS' if not errors else 'FAIL', 'output': str(Path(path).resolve()),
              'source_sha256': view['source_sha256'], 'result_fingerprint': view['result_fingerprint'],
              'delivery_sha256': content_hash, 'workbook_sha256': hashlib.sha256(Path(path).read_bytes()).hexdigest(),
              'item_count': len(model['items']), 'quantity_row_count': len(model['rows']),
              'missing_group_count': len(model['delivery'].get('missing_groups', [])),
              'business_choice_count': len(model['delivery'].get('decision_targets', [])),
              'current_null_count': sum(x.get('quantity') is None for x in model['items']),
              'states': {s: sum(x['state'] == s for x in model['rows']) for s in STATE_LABELS},
              'hyperlinks_checked': hyperlink_count, 'same_source_html': html_check,
              'row_checks': checks, 'errors': errors,
              'quantity_changes': 0, 'formula_count': 0,
              'validation_scope': '保存后逐项核对数值、空值、单位、成员、完整记录、长依据及内链；不代表重新量算或工程正式核准',
              'native_excel_validation': '未启动Excel；无公式，无原生重算要求'}
    wb.close()
    if errors:
        raise ValueError('工作簿核对失败：' + '；'.join(errors[:10]))
    return report


def build(input_path, output, report_path=None, html_path=None):
    input_path, output = Path(input_path).resolve(), Path(output).resolve()
    report_path = Path(report_path).resolve() if report_path else output.with_suffix('.核对回执.json')
    html_path = Path(html_path).resolve() if html_path else None
    protected = {input_path} | ({html_path} if html_path else set())
    if output in protected or report_path in protected or report_path == output:
        raise ValueError('成果或回执路径不能覆盖输入、HTML或彼此')
    if output.exists() or report_path.exists():
        raise ValueError('目标工作簿或核对回执已存在；请使用新版本名')
    input_bytes = input_path.read_bytes()
    input_sha256 = hashlib.sha256(input_bytes).hexdigest()
    raw = json.loads(input_bytes.decode('utf-8-sig'))
    del input_bytes
    if 'delivery' not in raw:
        raise ValueError('此Excel入口要求正式delivery视图；旧视图请先运行convert_qto_review_delivery.py --help，不能把旧技术合计冒充当前采用量')
    view = normalize_view(raw, input_path.parent)
    del raw
    wb, model, positions = create_workbook(view)
    output.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    # Verify a temporary file before publishing either artifact. x-mode protects
    # an already existing user file if another process writes during generation.
    with tempfile.TemporaryDirectory(prefix='qto-excel-', dir=output.parent) as temp:
        candidate = Path(temp) / 'candidate.xlsx'
        wb.save(candidate)
        wb.close()
        report = verify_workbook(candidate, view, model, positions, html_path)
        report['output'] = str(output)
        report['input'] = str(input_path)
        if hashlib.sha256(input_path.read_bytes()).hexdigest() != input_sha256:
            raise ValueError('生成期间来源输入发生变化，拒绝输出混合版本；请重新运行')
        report['input_sha256'] = input_sha256
        with output.open('xb') as stream:
            stream.write(candidate.read_bytes())
        with report_path.open('x', encoding='utf-8', newline='\n') as stream:
            stream.write(json.dumps(report, ensure_ascii=False, indent=2))
    return report


def main():
    parser = argparse.ArgumentParser(description='按需生成同源工程量汇总、明细两表Excel；不重算专业工程量')
    parser.add_argument('--input', type=Path, required=True, help='共享HTML同一正式delivery view JSON')
    parser.add_argument('--output', type=Path, required=True, help='新建xlsx，拒绝覆盖')
    parser.add_argument('--report', type=Path, help='逐项核对回执，默认工作簿同名.核对回执.json')
    parser.add_argument('--html', type=Path, help='可选：核对已有共享HTML内嵌数据同源性，不自动生成HTML')
    args = parser.parse_args()
    try:
        report = build(args.input, args.output, args.report, args.html)
        # Complete row-level business data belongs in the receipt, not stdout.
        print(json.dumps({k: v for k, v in report.items() if k != 'row_checks'}, ensure_ascii=False, indent=2))
        return 0
    except (OSError, ValueError, TypeError, KeyError, AssertionError) as error:
        print(json.dumps({'status': 'FAIL', 'error': str(error)}, ensure_ascii=False))
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
